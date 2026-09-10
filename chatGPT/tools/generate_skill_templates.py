#!/usr/bin/env python3
"""產生 Office Workflow skill 用的 xlsx 範本與示範素材圖。

需要 openpyxl 與 pillow：
    python3 -m venv .venv && .venv/bin/pip install openpyxl pillow
    .venv/bin/python tools/generate_skill_templates.py
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw

BASE = Path(__file__).resolve().parent.parent
EX = BASE / "Skills/Examples"
HEAD = PatternFill("solid", fgColor="4A3428")
HF = Font(color="FFFFFF", bold=True)
THIN = Border(*[Side(style="thin", color="CCCCCC")] * 4)


def finish(ws, widths=None):
    for c in ws[1]:
        if c.value:
            c.fill, c.font = HEAD, HF
            c.alignment = Alignment(horizontal="center", vertical="center")
    for col in ws.columns:
        w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 4, 40)


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"  {path.relative_to(BASE)}")


# 03 報價單範本
def quotation():
    wb = Workbook(); ws = wb.active; ws.title = "報價單"
    ws["A1"] = "潮汐物流股份有限公司　服務報價單"
    ws["A1"].font = Font(size=16, bold=True)
    meta = [("報價單號", "{{QUOTE_NO}}"), ("報價日期", "{{QUOTE_DATE}}"),
            ("有效期限", "{{VALID_UNTIL}}"), ("客戶名稱", "{{CUSTOMER}}"),
            ("聯絡人", "{{CONTACT}}"), ("聯絡方式", "{{CONTACT_INFO}}")]
    for i, (k, v) in enumerate(meta, start=3):
        ws[f"A{i}"], ws[f"B{i}"] = k, v
        ws[f"A{i}"].font = Font(bold=True)
    r = 10
    for i, h in enumerate(["項次", "服務項目", "規格說明", "數量", "單位", "單價(TWD)", "小計(TWD)"], start=1):
        c = ws.cell(r, i, h); c.fill, c.font = HEAD, HF
        c.alignment = Alignment(horizontal="center")
    for n in range(1, 9):
        ws.cell(r + n, 1, n)
        for i in range(1, 8):
            ws.cell(r + n, i).border = THIN
    b = r + 9
    for lbl, key in [("小計", "{{SUBTOTAL}}"), ("折讓", "{{DISCOUNT}}"), ("總計", "{{TOTAL}}")]:
        ws.cell(b, 6, lbl).font = Font(bold=True)
        ws.cell(b, 7, key); b += 1
    ws.cell(b + 1, 1, "交易條款").font = Font(bold=True)
    for i, t in enumerate(["付款方式：{{PAYMENT_TERMS}}", "交付方式：{{DELIVERY}}",
                           "備註：{{NOTES}}"], start=b + 2):
        ws.cell(i, 1, t)
    ws.cell(b + 6, 1, "※ 佔位符 {{...}} 由 skill 填入。未提供的資訊須向使用者確認，不得自行填入。")
    finish(ws)
    save(wb, EX / "Office_Workflow_03_Customer_Quotation/templates/quotation-template.xlsx")


# 04 行動計畫範本
def action_plan():
    wb = Workbook(); ws = wb.active; ws.title = "行動計畫"
    ws["A1"] = "會議行動計畫　{{MEETING_NAME}}"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = "會議日期：{{MEETING_DATE}}　　記錄：{{RECORDER}}"
    for i, h in enumerate(["編號", "行動項目", "負責人", "期限", "狀態", "相依於", "出處", "備註"], start=1):
        c = ws.cell(4, i, h); c.fill, c.font = HEAD, HF
    for n in range(1, 16):
        ws.cell(4 + n, 1, n)
        for i in range(1, 9):
            ws.cell(4 + n, i).border = THIN
    ws.cell(22, 1, "狀態值：未開始／進行中／已完成／待確認").font = Font(bold=True)
    ws.cell(23, 1, "※ 負責人或期限缺一，狀態一律填「待確認」，不可自行推估。")
    ws2 = wb.create_sheet("待確認事項")
    for i, h in enumerate(["編號", "事項", "缺什麼", "該問誰", "出處"], start=1):
        c = ws2.cell(1, i, h); c.fill, c.font = HEAD, HF
    finish(ws); finish(ws2)
    save(wb, EX / "Office_Workflow_04_Meeting_Action_Plan/templates/action-plan-template.xlsx")


# 05 交接清單範本
def handover():
    wb = Workbook(); ws = wb.active; ws.title = "交接清單"
    ws["A1"] = "職務代理交接清單"
    ws["A1"].font = Font(size=14, bold=True)
    for i, (k, v) in enumerate([("請假人", "{{APPLICANT}}"), ("職務", "{{ROLE}}"),
                                ("請假期間", "{{LEAVE_PERIOD}}"), ("代理人", "{{DEPUTY}}"),
                                ("交接日期", "{{HANDOVER_DATE}}")], start=3):
        ws[f"A{i}"], ws[f"B{i}"] = k, v
        ws[f"A{i}"].font = Font(bold=True)
    for i, h in enumerate(["編號", "交接事項", "現況", "期間內需處理", "急迫度", "聯絡窗口", "備註"], start=1):
        c = ws.cell(9, i, h); c.fill, c.font = HEAD, HF
    for n in range(1, 13):
        ws.cell(9 + n, 1, n)
        for i in range(1, 8):
            ws.cell(9 + n, i).border = THIN
    ws.cell(23, 1, "急迫度：🔴 期間內必須處理／🟡 可延後／🟢 回來再處理").font = Font(bold=True)
    ws2 = wb.create_sheet("代理權限")
    for i, h in enumerate(["項目", "代理人可否決行", "上限", "需請示對象"], start=1):
        c = ws2.cell(1, i, h); c.fill, c.font = HEAD, HF
    for r, row in enumerate([["日常請款", "可", "NT$10,000", "—"],
                             ["採購核決", "否", "—", "部門主管"],
                             ["對外正式函文", "否", "—", "部門主管"],
                             ["客訴回覆", "可（草稿）", "—", "客服主管確認後發出"]], start=2):
        for i, v in enumerate(row, start=1):
            ws2.cell(r, i, v)
    finish(ws); finish(ws2)
    save(wb, EX / "Office_Workflow_05_Leave_Handover/templates/leave-handover-template.xlsx")


# 06 採購比價範本
def purchase():
    wb = Workbook(); ws = wb.active; ws.title = "採購比價"
    ws["A1"] = "採購比價表　{{ITEM_NAME}}"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = "申請單位：{{DEPT}}　申請日期：{{DATE}}　需求數量：{{QTY}}"
    for i, h in enumerate(["廠商", "型號規格", "單價(TWD)", "數量", "小計", "交期(天)",
                           "保固", "付款條件", "資料來源", "報價日期"], start=1):
        c = ws.cell(4, i, h); c.fill, c.font = HEAD, HF
    for n in range(1, 6):
        for i in range(1, 11):
            ws.cell(4 + n, i).border = THIN
    ws.cell(11, 1, "比價結論").font = Font(bold=True)
    for i, t in enumerate(["建議廠商：{{RECOMMEND}}", "理由：{{REASON}}",
                           "※ 若最低價未被推薦，必須說明原因（交期、保固、資格）。",
                           "※ 每筆報價須註明資料來源與日期，來源不明者不得列入比較。"], start=12):
        ws.cell(i, 1, t)
    finish(ws)
    save(wb, EX / "Office_Workflow_06_Purchase_Checker/templates/purchase-compare-template.xlsx")


# 07 CRM 範本
def crm():
    wb = Workbook(); ws = wb.active; ws.title = "客戶互動紀錄"
    for i, h in enumerate(["日期", "客戶", "聯絡人", "管道", "類型", "摘要",
                           "我方承諾", "承諾期限", "後續動作", "負責人", "狀態"], start=1):
        c = ws.cell(1, i, h); c.fill, c.font = HEAD, HF
    for n in range(2, 22):
        for i in range(1, 12):
            ws.cell(n, i).border = THIN
    ws2 = wb.create_sheet("承諾追蹤")
    for i, h in enumerate(["承諾事項", "客戶", "承諾日", "期限", "是否兌現", "逾期天數", "風險"], start=1):
        c = ws2.cell(1, i, h); c.fill, c.font = HEAD, HF
    ws2.cell(12, 1, "※ 「我方承諾」欄必須引用原文，不可改寫或美化。").font = Font(bold=True)
    ws2.cell(13, 1, "※ 語意模糊的表述（「我們盡量」）標示為「需確認是否構成承諾」。")
    ws3 = wb.create_sheet("欄位說明")
    for r, row in enumerate([["管道", "電話／Email／面談／Slack／其他"],
                             ["類型", "詢價／客訴／需求變更／例行聯繫／合約"],
                             ["狀態", "待處理／處理中／已結案／待客戶回覆"]], start=1):
        ws3.cell(r, 1, row[0]).font = Font(bold=True)
        ws3.cell(r, 2, row[1])
    finish(ws); finish(ws2)
    save(wb, EX / "Office_Workflow_07_Customer_CRM/templates/crm-activity-template.xlsx")


# 示範素材圖（模擬掃描件）
def slips():
    specs = [
        (EX / "Office_Workflow_05_Leave_Handover/sample_materials/01_請假單.png",
         "TideFlow Logistics - Leave Request", [
             "Employee: Huang Chun-An (E-2041)", "Dept: Facilities Div.",
             "Leave Type: Annual Leave", "Period: 2026-09-21 to 2026-09-30",
             "Working Days: 8", "Deputy: Lin Wei-Lun",
             "Reason: Family trip", "Status: APPROVED 2026-09-05"]),
        (EX / "Office_Workflow_05_Leave_Handover/sample_materials/02_代理授權書.png",
         "Deputy Authorization", [
             "Principal: Huang Chun-An", "Deputy: Lin Wei-Lun",
             "Period: 2026-09-21 to 2026-09-30",
             "Authorized: routine payment <= NT$10,000",
             "NOT authorized: purchase approval",
             "NOT authorized: official external letters",
             "Approved by: Dept. Manager 2026-09-05"]),
        (EX / "Office_Workflow_06_Purchase_Checker/sample_materials/01_廠商報價單.png",
         "LengYan Equipment Co. - Quotation", [
             "Quote No: LY-2026-0912", "Date: 2026-09-12", "Valid: 30 days",
             "Item: Cold-chain container 60L", "Unit Price: NT$1,250",
             "Qty: 200", "Subtotal: NT$250,000",
             "Lead time: 21 days", "Warranty: 12 months",
             "Payment: 30 days after delivery"]),
        (EX / "Office_Workflow_06_Purchase_Checker/sample_materials/02_電商比價截圖.png",
         "Online Marketplace - Price Comparison", [
             "Cold-chain container 60L (equivalent spec)",
             "Seller A: NT$1,180  | ship 7d  | warranty 6m",
             "Seller B: NT$1,090  | ship 14d | warranty NONE",
             "Seller C: NT$1,320  | ship 3d  | warranty 12m",
             "", "Note: specs not verified as identical",
             "Captured: 2026-09-13"]),
        (EX / "Office_Workflow_07_Customer_CRM/sample_materials/01_客訴郵件.png",
         "Email - Damage Claim (3rd follow-up)", [
             "From: Chang Chih-Ming <dinghe-trading>",
             "Date: 2026-09-02 16:45",
             "Order: TFL-2026-081203",
             "Claim amount: NT$28,400",
             "First submitted: 2026-08-19",
             "Promised reply: within 7 working days",
             "Actual: no reply after 2 weeks",
             "Customer states: may terminate contract"]),
        (EX / "Office_Workflow_07_Customer_CRM/sample_materials/02_業務擴展詢問.png",
         "Email - Cold Chain Inquiry", [
             "From: Wang Shih-Han <maiho-bakery>",
             "Date: 2026-09-02 14:20",
             "Need: 8 stores, daily 03:00-06:00",
             "Temperature: 0-4 C, no break",
             "Start: 2026-10-05", "Volume: ~600 boxes/day",
             "Question: add-on to existing contract",
             "         or new contract?"]),
    ]
    for path, title, lines in specs:
        path.parent.mkdir(parents=True, exist_ok=True)
        W, H = 900, 560
        img = Image.new("RGB", (W, H), (252, 251, 248))
        d = ImageDraw.Draw(img)
        d.rectangle([20, 20, W - 20, H - 20], outline=(120, 110, 100), width=2)
        d.rectangle([20, 20, W - 20, 78], fill=(74, 52, 40))
        d.text((44, 44), title, fill=(255, 255, 255))
        y = 120
        for ln in lines:
            d.text((56, y), ln, fill=(40, 35, 30))
            y += 42
        d.text((56, H - 56), "* Fictional teaching material. Not a real document.",
               fill=(150, 140, 130))
        img.save(path)
        print(f"  {path.relative_to(BASE)}")


if __name__ == "__main__":
    print("產生 skill 範本與素材：")
    quotation(); action_plan(); handover(); purchase(); crm(); slips()
    print("\n完成。")
