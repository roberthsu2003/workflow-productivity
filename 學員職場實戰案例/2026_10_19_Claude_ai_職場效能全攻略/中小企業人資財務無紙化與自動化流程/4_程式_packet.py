# -*- coding: utf-8 -*-
"""無紙化流程｜合併列印工具
功能：
  1. 由「簽核歷程」Excel 產生簽核歷程頁
  2. 將傳票、請款單、簽核歷程、憑證附件依序合併成單一 PDF
  3. 自動加上目錄頁、PDF 書籤、每頁頁尾（單據編號與頁碼）

用法（範例）：
  python packet.py --doc-id EXP-202610-001 --title "費用報銷歸檔資料" ^
      --trail 無紙化流程_素材總表.xlsx ^
      --section "會計傳票=voucher.pdf" --section "已簽核請款單=claim.pdf" ^
      --section "簽核歷程=@trail" --section "發票／憑證附件=receipt.pdf" ^
      --out EXP-202610-001_歸檔.pdf

  "@trail" 代表由 Excel 簽核歷程自動產生該頁。
字型：預設使用 Windows 微軟正黑體；其他系統請用 --font 指定 .ttc/.ttf 路徑。
"""
import argparse, io, os, sys, datetime
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.pdfgen import canvas

FONT_CANDIDATES = [
    r"C:\Windows\Fonts\msjh.ttc",
    "/System/Library/Fonts/PingFang.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
]
FN = "CJK"

def register_font(path=None):
    for p in ([path] if path else FONT_CANDIDATES):
        if p and os.path.exists(p):
            for idx in range(4):
                try:
                    pdfmetrics.registerFont(TTFont(FN, p, subfontIndex=idx))
                    return p
                except Exception:
                    continue
    sys.exit("找不到可用的中文字型，請用 --font 指定 .ttf/.ttc 路徑")

def styles():
    return dict(
        h1=ParagraphStyle("h1", fontName=FN, fontSize=18, leading=24, textColor=colors.HexColor("#1F4E79"), spaceAfter=6),
        sub=ParagraphStyle("sub", fontName=FN, fontSize=9, leading=13, textColor=colors.HexColor("#7F7F7F")),
        body=ParagraphStyle("body", fontName=FN, fontSize=10, leading=15),
        cell=ParagraphStyle("cell", fontName=FN, fontSize=9, leading=13),
    )

def grid(data, widths, header=True):
    st = styles()["cell"]
    rows = [[Paragraph(str(c), st) for c in r] for r in data]
    t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
    cmds = [("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#BFBFBF")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
    if header:
        cmds.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DEEAF6")))
    t.setStyle(TableStyle(cmds))
    return t

def build(story, path):
    SimpleDocTemplate(path, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm,
                      topMargin=18*mm, bottomMargin=20*mm).build(story)

def trail_pdf(xlsx, doc_id, out, sheet="簽核歷程格式"):
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, data_only=True)[sheet]
    rows = [["關卡", "簽核人", "動作", "意見", "時間"]]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[0] == doc_id:
            rows.append([r[2] or "", r[3] or "", r[4] or "", r[5] or "", str(r[6] or "")])
    if len(rows) == 1:
        sys.exit(f"簽核歷程中找不到單據編號：{doc_id}")
    s = styles()
    story = [Paragraph("簽核歷程", s["h1"]), Paragraph(f"單據編號：{doc_id}", s["sub"]), Spacer(1, 8),
             grid(rows, [34*mm, 30*mm, 22*mm, 50*mm, 38*mm])]
    build(story, out)

def count_pages(p):
    return len(PdfReader(p).pages)

def toc_pdf(title, doc_id, items, out):
    s = styles()
    data = [["項次", "內容", "起始頁"]] + [[str(i), n, str(pg)] for i, (n, pg) in enumerate(items, 1)]
    story = [Paragraph(title, s["h1"]),
             Paragraph(f"單據編號：{doc_id}　產生日期：{datetime.date.today():%Y-%m-%d}", s["sub"]),
             Spacer(1, 10), Paragraph("目錄", s["body"]), Spacer(1, 4),
             grid(data, [18*mm, 110*mm, 28*mm])]
    build(story, out)

def stamp(reader_pages, doc_id):
    total = len(reader_pages)
    for i, page in enumerate(reader_pages, 1):
        w = float(page.mediabox.width); h = float(page.mediabox.height)
        buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=(w, h))
        c.setFont(FN, 8); c.setFillColor(colors.HexColor("#7F7F7F"))
        c.drawCentredString(w/2, 8*mm, f"{doc_id}　第 {i} / {total} 頁")
        c.save(); buf.seek(0)
        page.merge_page(PdfReader(buf).pages[0])

def merge(doc_id, title, sections, out, trail_xlsx=None, tmp="_tmp_packet"):
    os.makedirs(tmp, exist_ok=True)
    files = []
    for name, src in sections:
        if src == "@trail":
            if not trail_xlsx: sys.exit("使用 @trail 需加上 --trail Excel 路徑")
            p = os.path.join(tmp, "trail.pdf"); trail_pdf(trail_xlsx, doc_id, p); src = p
        if not os.path.exists(src): sys.exit(f"找不到檔案：{src}")
        files.append((name, src, count_pages(src)))
    toc_len = 1
    for _ in range(2):  # 目錄頁數可能影響頁碼，算兩次
        start, items = toc_len + 1, []
        for name, _, n in files:
            items.append((name, start)); start += n
        tp = os.path.join(tmp, "toc.pdf"); toc_pdf(title, doc_id, items, tp)
        toc_len = count_pages(tp)
    w = PdfWriter()
    for pg in PdfReader(tp).pages: w.add_page(pg)
    w.add_outline_item("目錄", 0)
    cur = toc_len
    for name, src, n in files:
        for pg in PdfReader(src).pages: w.add_page(pg)
        w.add_outline_item(name, cur); cur += n
    stamp(w.pages, doc_id)
    w.add_metadata({"/Title": f"{title} {doc_id}"})
    with open(out, "wb") as f: w.write(f)
    return out

if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="無紙化流程合併列印工具")
    ap.add_argument("--doc-id", required=True); ap.add_argument("--title", default="歸檔資料")
    ap.add_argument("--section", action="append", required=True, help='格式："名稱=檔案.pdf" 或 "名稱=@trail"')
    ap.add_argument("--trail"); ap.add_argument("--out", required=True); ap.add_argument("--font")
    a = ap.parse_args(); register_font(a.font)
    secs = [tuple(s.split("=", 1)) for s in a.section]
    print("完成：", merge(a.doc_id, a.title, secs, a.out, a.trail))
